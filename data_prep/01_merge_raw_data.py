"""
01_merge_raw_data.py

Loads all 24 raw KMC line-listing files (12 Dengue + 12 Malaria, 2014-2025),
harmonizes their inconsistent column naming across years into a common
schema, and merges them into one master DataFrame.

Target harmonized columns (per Ankur's brief):
    Date, Month, Age, Sex, Ward, Borough, Treatment_Raw, Disease, Year

Personally identifying columns (Name, Address, Contact No, Premises No,
Religion, hospital contact person) are dropped by design - they are not
needed for the RQs and should not be embedded in a shared research dataset.

Run: python3 01_merge_raw_data.py
"""
import os
import re
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.environ.get("RAW_DIR") or os.path.join(HERE, "..", "data", "raw")   # the 24 KMC workbooks (not included)
DENGUE_DIR = os.path.join(BASE, "DENGUE - 2014 - 2025")
MALARIA_DIR = os.path.join(BASE, "MALARIA - 2014 - 2025")
OUT_DIR = os.environ.get("DATA_DIR") or os.path.join(HERE, "..", "data")
os.makedirs(OUT_DIR, exist_ok=True)

YEARS = list(range(2014, 2026))

# Header-row configuration derived from manual inspection of all 24 files
# (header_row is 1-indexed row containing the column names;
#  data_start_row is the 1-indexed row where actual case rows begin)
def header_config(disease, year):
    if disease == "Dengue":
        return {"header_row": 2, "data_start_row": 3}
    else:  # Malaria
        if year == 2014:
            return {"header_row": 1, "data_start_row": 3}   # header + PV/PF/MIXED subheader, no title banner
        elif year == 2015:
            return {"header_row": 2, "data_start_row": 3}   # title banner + header
        else:
            return {"header_row": 2, "data_start_row": 4}   # title banner + header + PV/PF/MIXED subheader


def normalize(h):
    if h is None:
        return ""
    h = str(h).replace("\n", " ")
    h = re.sub(r"\s+", " ", h).strip().lower()
    h = h.rstrip(".")
    return h


def find_target_columns(headers):
    """headers: list of raw header values (in column order). Returns dict field->index."""
    norm = [normalize(h) for h in headers]
    result = {}

    # Date
    date_idx = None
    for i, h in enumerate(norm):
        if "reporting date" in h:
            date_idx = i
            break
    if date_idx is None:
        for i, h in enumerate(norm):
            if "blood test" in h:  # covers "blood test date" / "blood testing date"
                date_idx = i
                break
    if date_idx is None:
        for i, h in enumerate(norm):
            if h == "date":
                date_idx = i
                break
    result["Date"] = date_idx

    # Month
    result["Month"] = next((i for i, h in enumerate(norm) if h == "month"), None)

    # Age
    result["Age"] = next((i for i, h in enumerate(norm) if "age" in h), None)

    # Sex
    result["Sex"] = next((i for i, h in enumerate(norm) if "sex" in h), None)

    # Ward
    result["Ward"] = next((i for i, h in enumerate(norm) if h == "ward"), None)

    # Borough / Br
    result["Borough"] = next((i for i, h in enumerate(norm) if h in ("borough", "br")), None)

    # Treatment proxy (place of admission / home-treatment / hospital treatment / nursing home)
    treat_idx = None
    for i, h in enumerate(norm):
        if "date" in h:
            continue
        if "dmission" in h or ("home" in h and "treatment" in h) or ("hospital" in h and "treatment" in h) or "nursing home" in h:
            treat_idx = i
            break
    result["Treatment_Raw"] = treat_idx

    return result


def load_sheet_rows(ws, header_row, data_start_row, max_col=40):
    """Return (header_values, data_rows) for a worksheet."""
    header_values = None
    for i, row in enumerate(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=max_col, values_only=True), start=header_row):
        header_values = list(row)
    data_rows = []
    for row in ws.iter_rows(min_row=data_start_row, max_col=max_col, values_only=True):
        if any(v is not None and str(v).strip() != "" for v in row):
            data_rows.append(row)
    return header_values, data_rows


def pick_best_sheet(wb, header_row):
    """When a file has multiple sheets, pick the one whose header row matches the most target fields."""
    best_name, best_score, best_headers = None, -1, None
    for name in wb.sheetnames:
        ws = wb[name]
        try:
            headers = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
        except StopIteration:
            continue
        mapping = find_target_columns(headers)
        score = sum(1 for v in mapping.values() if v is not None)
        if score > best_score:
            best_name, best_score, best_headers = name, score, headers
    return best_name


def process_file(path, disease, year):
    cfg = header_config(disease, year)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if len(wb.sheetnames) == 1:
        sheet_name = wb.sheetnames[0]
    else:
        sheet_name = pick_best_sheet(wb, cfg["header_row"])
    ws = wb[sheet_name]
    headers, data_rows = load_sheet_rows(ws, cfg["header_row"], cfg["data_start_row"])
    wb.close()

    mapping = find_target_columns(headers)
    records = []
    for row in data_rows:
        rec = {
            "Date_Raw": row[mapping["Date"]] if mapping["Date"] is not None and mapping["Date"] < len(row) else None,
            "Month_Raw": row[mapping["Month"]] if mapping["Month"] is not None and mapping["Month"] < len(row) else None,
            "Age_Raw": row[mapping["Age"]] if mapping["Age"] is not None and mapping["Age"] < len(row) else None,
            "Sex_Raw": row[mapping["Sex"]] if mapping["Sex"] is not None and mapping["Sex"] < len(row) else None,
            "Ward_Raw": row[mapping["Ward"]] if mapping["Ward"] is not None and mapping["Ward"] < len(row) else None,
            "Borough_Raw": row[mapping["Borough"]] if mapping["Borough"] is not None and mapping["Borough"] < len(row) else None,
            "Treatment_Raw": row[mapping["Treatment_Raw"]] if mapping["Treatment_Raw"] is not None and mapping["Treatment_Raw"] < len(row) else None,
            "Disease": disease,
            "Year": year,
            "Source_Sheet": sheet_name,
        }
        records.append(rec)

    return records, mapping, sheet_name


def main():
    all_records = []
    audit_rows = []

    for year in YEARS:
        dengue_path = os.path.join(DENGUE_DIR, f"DENGUE LINE-LIST {year}.xlsx")
        records, mapping, sheet_name = process_file(dengue_path, "Dengue", year)
        all_records.extend(records)
        audit_rows.append({"Disease": "Dengue", "Year": year, "Sheet": sheet_name, "Rows": len(records), **{f"col_{k}": v for k, v in mapping.items()}})

        malaria_path = os.path.join(MALARIA_DIR, f"MALARIA {year}.xlsx")
        records, mapping, sheet_name = process_file(malaria_path, "Malaria", year)
        all_records.extend(records)
        audit_rows.append({"Disease": "Malaria", "Year": year, "Sheet": sheet_name, "Rows": len(records), **{f"col_{k}": v for k, v in mapping.items()}})

    df = pd.DataFrame(all_records)
    audit = pd.DataFrame(audit_rows)

    out_path = os.path.join(OUT_DIR, "merged_raw.csv")
    audit_path = os.path.join(OUT_DIR, "merge_audit.csv")
    df.to_csv(out_path, index=False)
    audit.to_csv(audit_path, index=False)

    print(f"Merged dataset: {len(df)} rows -> {out_path}")
    print(f"\nPer-file audit (row counts + which column index was matched for each field):")
    print(audit.to_string(index=False))
    print(f"\nDisease counts:\n{df['Disease'].value_counts()}")
    print(f"\nYear x Disease pivot (total rows, not just non-null Date):\n{df.pivot_table(index='Year', columns='Disease', values='Source_Sheet', aggfunc='size', fill_value=0)}")


if __name__ == "__main__":
    main()
